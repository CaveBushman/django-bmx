from io import BytesIO

from django.test import SimpleTestCase
from openpyxl import Workbook
from PIL import Image

from bmx.excel_utils import HEADER_FILL, HEADER_FONT, set_column_widths, write_styled_header
from bmx.image_utils import normalize_avatar_image


class NormalizeAvatarImageTests(SimpleTestCase):
    def _image_file(self, size=(800, 600), fmt="PNG"):
        buffer = BytesIO()
        Image.new("RGB", size, "red").save(buffer, format=fmt)
        buffer.seek(0)
        return buffer

    def test_returns_square_image_of_configured_size(self):
        with self.settings(AVATAR_FINAL_IMAGE_SIZE=128):
            content, extension = normalize_avatar_image(self._image_file())

        self.assertIn(extension, {"webp", "jpg"})
        result = Image.open(BytesIO(content))
        self.assertEqual(result.size, (128, 128))

    def test_invalid_image_raises_pillow_error(self):
        with self.assertRaises(OSError):
            normalize_avatar_image(BytesIO(b"not an image"))


class ExcelHeaderHelpersTests(SimpleTestCase):
    def test_write_styled_header_styles_every_header_cell(self):
        ws = Workbook().active
        write_styled_header(ws, ["A", "B", "C"])

        self.assertEqual([cell.value for cell in ws[1]], ["A", "B", "C"])
        for cell in ws[1]:
            self.assertEqual(cell.font.b, HEADER_FONT.b)
            self.assertEqual(cell.fill.fgColor.rgb, HEADER_FILL.fgColor.rgb)

    def test_set_column_widths(self):
        ws = Workbook().active
        set_column_widths(ws, [10, 20])

        self.assertEqual(ws.column_dimensions["A"].width, 10)
        self.assertEqual(ws.column_dimensions["B"].width, 20)
